from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.shortcuts import get_object_or_404, redirect, render

from .forms import MobileForm
from .models import Mobile

IMPORT_FIELDS = ('brand', 'model_name', 'ram', 'storage', 'price', 'condition', 'description')


@login_required
def home(request):
    return render(request, 'mobiles/home.html', {'mobiles': Mobile.objects.order_by('-created_at')})


@login_required
def add_mobile(request):
    if request.method == 'POST':
        form = MobileForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            messages.success(request, 'Mobile added to the inventory.')
            return redirect('home')
    else:
        form = MobileForm()
    return render(request, 'mobiles/add_mobile.html', {'form': form})


@login_required
def import_mobiles(request):
    if request.method == 'POST':
        uploaded = request.FILES.get('excel_file')
        if not uploaded or not uploaded.name.lower().endswith('.xlsx'):
            messages.error(request, 'Choose an Excel .xlsx file to continue.')
            return redirect('import_mobiles')
        try:
            from openpyxl import load_workbook

            sheet = load_workbook(uploaded, read_only=True, data_only=True).active
            data = list(sheet.iter_rows(values_only=True))
            if not data:
                raise ValueError('The workbook is empty.')
            headers = {str(value).strip().lower(): index for index, value in enumerate(data[0]) if value is not None}
            missing = [field for field in IMPORT_FIELDS[:-1] if field not in headers]
            if missing:
                raise ValueError('Missing required columns: ' + ', '.join(missing))
        except Exception as error:
            messages.error(request, f'Could not read the workbook: {error}')
            return redirect('import_mobiles')
        rows, errors = [], []
        for row_number, values in enumerate(data[1:], 2):
            item = {field: str(values[headers[field]] or '').strip() if headers.get(field) is not None and len(values) > headers[field] else '' for field in IMPORT_FIELDS}
            if not any(item.values()):
                continue
            form = MobileForm(item)
            problems = []
            if not form.is_valid():
                problems.extend(
                    f'{field.replace("_", " ").title()}: {message}'
                    for field, messages_list in form.errors.items()
                    for message in messages_list
                )
            if problems:
                errors.append({'row': row_number, 'message': '; '.join(problems)})
            else:
                rows.append({field: str(form.cleaned_data[field]) for field in IMPORT_FIELDS})
        request.session['mobile_import_rows'] = rows
        return render(request, 'mobiles/import_review.html', {'rows': rows, 'errors': errors})
    return render(request, 'mobiles/import_mobiles.html', {'fields': IMPORT_FIELDS})


@login_required
def confirm_import(request):
    rows = request.session.get('mobile_import_rows', [])
    if request.method == 'POST' and rows:
        Mobile.objects.bulk_create([Mobile(**row) for row in rows])
        del request.session['mobile_import_rows']
        messages.success(request, f'{len(rows)} mobile record(s) uploaded successfully.')
        return redirect('home')
    messages.warning(request, 'There is no reviewed import ready to upload.')
    return redirect('import_mobiles')


@login_required
def edit_mobile(request, id):
    mobile = get_object_or_404(Mobile, id=id)
    if request.method == 'POST':
        form = MobileForm(request.POST, request.FILES, instance=mobile)
        if form.is_valid():
            form.save()
            messages.success(request, 'Mobile details updated.')
            return redirect('home')
    else:
        form = MobileForm(instance=mobile)
    return render(request, 'mobiles/edit_mobile.html', {'mobile': mobile, 'form': form})


@login_required
def delete_mobile(request, id):
    mobile = get_object_or_404(Mobile, id=id)
    if request.method == 'POST':
        mobile.delete()
        messages.success(request, 'Mobile removed from inventory.')
        return redirect('home')
    return render(request, 'mobiles/delete_mobile.html', {'mobile': mobile})
